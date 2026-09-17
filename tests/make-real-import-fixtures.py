"""Generate full Office packages with maintained authoring libraries, then convert to ODF.
All personal details are fictional. No Office/LibreOffice account is used.
Run with a Python containing python-pptx, python-docx and openpyxl.
LibreOffice conversion uses an isolated temporary profile and never the user's profile.
"""
from pathlib import Path
import argparse, hashlib, importlib.metadata, json, shutil, subprocess, tempfile
from datetime import datetime
from pptx import Presentation
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
ap=argparse.ArgumentParser();ap.add_argument('output',nargs='?',default=str(Path(__file__).parent/'fixtures/imports'));ap.add_argument('--soffice',default=shutil.which('soffice'));args=ap.parse_args()
out=Path(args.output);out.mkdir(parents=True,exist_ok=True)
manifest_path=out/'fixture-manifest.json';manifest=json.loads(manifest_path.read_text())
created=[]
def add(name,contains,**other):
 p=out/name;created.append({'file':name,'contains':contains,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'fixtureKind':'library-generated-full-package',**other})
# Ordinary complete PowerPoint with themes, layout and master parts from python-pptx.
pres=Presentation();pres.core_properties.author='Fictional QA Author';pres.core_properties.title='Fictional staff privacy fixture'
for i,name in enumerate(['FIRST_GENERATED_SLIDE: Sofia Marin','SECOND_GENERATED_SLIDE: Amir Haddad','HIDDEN_GENERATED_SLIDE: Noor Hassan']):
 s=pres.slides.add_slide(pres.slide_layouts[1]);s.shapes.title.text=name
 frame=s.placeholders[1].text_frame;frame.clear();para=frame.paragraphs[0]
 para.add_run().text='Birth';para.add_run().text='date: 14/03/2010'
 if i==0:s.notes_slide.notes_text_frame.text='GENERATED_NOTES: Aya Hassan. notes.only@example.org'
 if i==2:s._element.set('show','0')
pres.save(out/'generated-staff-presentation.pptx')
assert len(Presentation(out/'generated-staff-presentation.pptx').slides)==3
add('generated-staff-presentation.pptx',['Birthdate: 14/03/2010','GENERATED_NOTES: Aya Hassan'],ordered=['FIRST_GENERATED_SLIDE','SECOND_GENERATED_SLIDE','HIDDEN_GENERATED_SLIDE'])
# Complete Word package with header, footer, split runs and a comment.
doc=Document();doc.core_properties.author='Fictional QA Author';doc.add_heading('GENERATED_WORD_BODY: Sofia Marin',0)
p=doc.add_paragraph();runs=[p.add_run('Birth'),p.add_run('date: 14/03/2010')]
doc.sections[0].header.paragraphs[0].text='GENERATED_WORD_HEADER: Amir Haddad'
doc.sections[0].footer.paragraphs[0].text='GENERATED_WORD_FOOTER: footer.only@example.org'
assert hasattr(doc,'add_comment'),'python-docx with comment support required'
doc.add_comment(runs,text='GENERATED_WORD_COMMENT: private.note@example.org',author='Fictional Reviewer',initials='FR')
doc.save(out/'generated-staff-document.docx');assert Document(out/'generated-staff-document.docx').paragraphs
add('generated-staff-document.docx',['GENERATED_WORD_BODY','GENERATED_WORD_HEADER','GENERATED_WORD_FOOTER','GENERATED_WORD_COMMENT','Birthdate: 14/03/2010'])
# Complete Excel package, styled DOB cell, comment and hidden sheet.
wb=Workbook();ws=wb.active;ws.title='Fictional register';ws.append(['Name','Date of birth','Contact'])
ws.append(['Sofia Marin',datetime(2010,3,14),'sofia@example.org']);ws['B2'].number_format='dd/mm/yyyy'
ws['A2'].comment=Comment('GENERATED_EXCEL_COMMENT: private.note@example.org','Fictional Reviewer')
hidden=wb.create_sheet('Private context');hidden['A1']='GENERATED_HIDDEN_SHEET: Amir Haddad';hidden.sheet_state='hidden'
wb.save(out/'generated-staff-workbook.xlsx');assert load_workbook(out/'generated-staff-workbook.xlsx')['Private context'].sheet_state=='hidden'
add('generated-staff-workbook.xlsx',['Sofia Marin','sofia@example.org','GENERATED_HIDDEN_SHEET','GENERATED_EXCEL_COMMENT','2010-03-14'])
# LibreOffice round trips these authored packages to real ODF packages.
version=None
if args.soffice:
 version=subprocess.run([args.soffice,'--version'],check=True,text=True,capture_output=True,timeout=30).stdout.strip()
 with tempfile.TemporaryDirectory(prefix='privacy-fixture-lo-') as profile:
  for source,ext,contains in [('generated-staff-presentation.pptx','odp',['FIRST_GENERATED_SLIDE','GENERATED_NOTES','HIDDEN_GENERATED_SLIDE']),('generated-staff-document.docx','odt',['GENERATED_WORD_BODY','GENERATED_WORD_HEADER','GENERATED_WORD_FOOTER','GENERATED_WORD_COMMENT']),('generated-staff-workbook.xlsx','ods',['Sofia Marin','sofia@example.org','GENERATED_HIDDEN_SHEET','GENERATED_EXCEL_COMMENT'])]:
   result=subprocess.run([args.soffice,'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to',ext,'--outdir',str(out),str(out/source)],check=True,text=True,capture_output=True,timeout=60)
   target=Path(source).stem+'.'+ext
   if not (out/target).is_file():raise RuntimeError('LibreOffice produced no '+target+': '+result.stdout+' '+result.stderr)
   add(target,contains,fixtureKind='libreoffice-converted-full-package')
else:raise RuntimeError('LibreOffice/soffice required for the three ODF fixture conversions')
new_names={i['file'] for i in created};manifest['fixtures']=[i for i in manifest['fixtures'] if i['file'] not in new_names]+created
manifest['realPackageGenerator']={'python-pptx':importlib.metadata.version('python-pptx'),'python-docx':importlib.metadata.version('python-docx'),'openpyxl':importlib.metadata.version('openpyxl'),'libreoffice':version,'scope':'Generated files reopened by their authoring libraries and converted by LibreOffice. No interactive Microsoft Office application verification.'}
manifest_path.write_text(json.dumps(manifest,indent=2)+'\n')
print(json.dumps({'created':len(created),'totalFixtures':len(manifest['fixtures']),'output':str(out),'provenance':manifest['realPackageGenerator']}))
